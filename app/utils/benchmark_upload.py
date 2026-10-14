"""Read Excel snapshots without executing formulas or retaining uploaded files."""
from datetime import datetime, timezone
from io import BytesIO
from zipfile import ZipFile

import openpyxl

from app.utils import client_benchmarks as b
from app.utils.benchmark_billing import read_roster

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_ROWS = 100_000
ACTIVITY_HEADERS = ('campaign id', 'client', 'month', 'company size', 'active', 'leads')
RAW_HEADERS = ('timestamp', 'client', 'lead company size', 'lead category')


def read_excel(content):
    if len(content) > MAX_UPLOAD_BYTES:
        raise b.SourceError('Upload an .xlsx file smaller than 20 MB.')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 200 * 1024 * 1024:
                raise b.SourceError('This workbook is too large when expanded. Export only the lead records tab.')
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
    except b.SourceError:
        raise
    except Exception as exc:
        raise b.SourceError('Could not read this workbook. Upload a valid, unencrypted .xlsx file.') from exc
    try:
        candidates = []
        for sheet in workbook:
            for index, row in enumerate(sheet.iter_rows(max_row=10, max_col=32, values_only=True)):
                names = [b._normal(cell) for cell in row]
                if all(h in names for h in ACTIVITY_HEADERS):
                    candidates.append((-1, sheet, index, names))
                    break
                if all(h in names for h in RAW_HEADERS):
                    candidates.append((0, sheet, index, names))
                    break
                if all(h in names for h in b.HEADERS):
                    candidates.append((1, sheet, index, names))
                    break
        if not candidates:
            raise b.SourceError('No lead data tab found. Include the eight supplied columns, or the original Form Responses tab.')
        priority = min(item[0] for item in candidates)
        matches = [item for item in candidates if item[0] == priority]
        if len(matches) != 1:
            raise b.SourceError('Multiple lead data tabs match. Upload a workbook with only the intended lead data tab.')
        kind, sheet, header_index, names = matches[0]
        required = ACTIVITY_HEADERS if kind == -1 else RAW_HEADERS if kind == 0 else b.HEADERS
        if any(names.count(h) != 1 for h in required):
            raise b.SourceError('Duplicate lead data headers found. Give each required column a unique name.')
        values = []
        for row in sheet.iter_rows(min_row=header_index+2, max_col=32, values_only=True):
            values.append(row)
            if len(values) > MAX_ROWS:
                raise b.SourceError('This tab exceeds 100,000 rows. Export a smaller reporting period.')
        if kind == -1:
            records, quality = parse_activity(names, values)
            definition = 'Campaign-month activity supplied in the workbook, including active months with zero leads. Leads are allocated to target-size bands; campaign-months are counted once per segment.'
        elif kind == 1:
            records, quality = b.parse_rows([names, *values])
            definition = 'Dated, named lead records from the supplied eight-column tab.'
        else:
            records, quality = parse_form_responses(names, values)
            definition = 'Only rows categorised Lead are counted. Interested and uncategorised records are excluded. A company name is not required.'
        if not records:
            raise b.SourceError('No valid lead records were found in the selected tab.')
        delivery_roster = read_roster(workbook) if kind != -1 else None
        dates = [r['date'] for r in records]
        return {('active_records' if kind == -1 else 'records'): records, 'quality': quality,
                'fetched_at': datetime.now(timezone.utc).isoformat(),
                'source_kind': 'excel', 'source_tab': sheet.title,
                **({'delivery_roster': delivery_roster} if delivery_roster is not None else {}),
                'date_min': min(dates).isoformat(), 'date_max': max(dates).isoformat(),
                'lead_definition': definition}
    finally:
        workbook.close()


def parse_form_responses(names, rows):
    records, seen = [], set()
    quality = {'invalid_rows': 0, 'duplicates': 0, 'unknown_size': 0, 'excluded_category': 0}
    indices = {h: names.index(h) for h in RAW_HEADERS}
    for row in rows:
        if not any(v is not None and str(v).strip() for v in row):
            continue
        def cell(name):
            i = indices[name]
            return row[i] if i < len(row) else None
        client = b._normal(cell('client'))
        day = b.parse_date(cell('timestamp'))
        if not client or not day:
            quality['invalid_rows'] += 1
            continue
        if b._normal(cell('lead category')) != 'lead':
            quality['excluded_category'] += 1
            continue
        signature = tuple(b._normal(v) for v in row)
        if signature in seen:
            quality['duplicates'] += 1
            continue
        seen.add(signature)
        band = b.band_index(cell('lead company size'))
        quality['unknown_size'] += band is None
        records.append({'client': client, 'date': day, 'band': band})
    return records, quality


def parse_activity(names, rows):
    """Require explicit activity, including zeros; never infer it from outcomes."""
    import re
    from datetime import date
    from decimal import Decimal, InvalidOperation
    indices = {h: names.index(h) for h in ACTIVITY_HEADERS}
    records, seen, owners = [], set(), {}
    quality = {'invalid_rows': 0, 'duplicates': 0, 'unknown_size': 0, 'inactive_rows': 0}
    for row in rows:
        if not any(v is not None and str(v).strip() for v in row):
            continue
        values = {h: row[i] if i < len(row) else None for h, i in indices.items()}
        campaign, client = (b._normal(values[h]) for h in ('campaign id', 'client'))
        raw_month = values['month']
        if isinstance(raw_month, (date, datetime)):
            raw_month = raw_month.strftime('%Y-%m')
        if not campaign or not client or not re.fullmatch(r'[1-9][0-9]{3}-(0[1-9]|1[0-2])', str(raw_month)):
            raise b.SourceError('Activity rows need a Campaign ID, Client and Month (YYYY-MM).')
        band = b.band_index(values['company size'])
        if band is None:
            raise b.SourceError('Every activity row needs one supported target company-size band.')
        active = str(values['active']).strip().lower()
        if active not in ('yes', 'no', 'true', 'false', '1', '0'):
            raise b.SourceError('Active must explicitly be Yes or No, including months with zero leads.')
        try:
            leads = Decimal(str(values['leads']))
            if not leads.is_finite() or leads < 0 or leads > 100000000 or leads != leads.to_integral_value():
                raise ValueError()
        except (InvalidOperation, ValueError):
            raise b.SourceError('Leads must be a non-negative whole number, including explicit zeroes.')
        if campaign in owners and owners[campaign] != client:
            raise b.SourceError('Each Campaign ID must belong to a single client.')
        owners[campaign] = client
        key = campaign, str(raw_month), band
        if key in seen:
            raise b.SourceError('Duplicate campaign/month/size rows found. Supply each combination once.')
        seen.add(key)
        if active in ('no', 'false', '0'):
            if leads:
                raise b.SourceError('An inactive campaign-month cannot have leads in this report. Reconcile activity and lead dates before importing.')
            quality['inactive_rows'] += 1
            continue
        records.append({'client': client, 'campaign': campaign, 'date': date.fromisoformat(str(raw_month)+'-01'), 'band': band, 'leads': int(leads)})
    return records, quality
