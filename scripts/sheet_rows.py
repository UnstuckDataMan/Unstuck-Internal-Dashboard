"""
Reads the old Credit Control Google Sheet (downloaded as .xlsx) into one plain row per invoice.

This is the same parsing the original July import used (header-driven, because the tab layout changed many times
between March 2022 and today), kept in one place so reset_from_sheet.py and any later check read the sheet the same way.

parse_sheet(path) -> (rows, warnings, month_tabs)
  rows         one dict per invoice line on a month tab: month, sheet, row, company, key, qb, email, am, pkg, terms, agreed,
               notes, mm, amount, amt_raw (the £ cell exactly as typed), usd, eur, issue, due, issued,
               paid (None = the tab has no Paid column), auto
  warnings     tabs that could not be read
  month_tabs   {month: tab name}
"""
from __future__ import annotations

import datetime
import re

import openpyxl

SKIP_SHEETS = {'-----', 'Retention', 'Team Commision', 'All clients', 'Sinarios'}
MONTH_WORDS = {m.lower(): i + 1 for i, m in enumerate(
    ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September', 'October', 'November', 'December'])}
for _i, _m in enumerate(['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']):
    MONTH_WORDS.setdefault(_m.lower(), _i + 1)

JUNK_COMPANY = {'client', 'clients', 'client company name', 'company name', 'total', 'no', 'yes',
                'latest month', 'issue automatically', 'issue manually', 'issue mannuallly', 'n/a', 'na', '-'}


def looks_junk(name) -> bool:
    s = str(name).strip()
    return s.lower() in JUNK_COMPANY or len(re.findall(r'[A-Za-z]', re.sub(r'\(.*?\)', '', s))) < 2


def tab_month_key(name: str):
    toks = name.strip().split()
    if not toks:
        return None
    m = MONTH_WORDS.get(toks[0].lower())
    if not m:
        return None
    year = 2022
    if len(toks) > 1:
        t = toks[-1]
        if re.fullmatch(r'\d{4}', t):
            year = int(t)
        elif re.fullmatch(r'\d{2}', t):
            year = 2000 + int(t)
    return f"{year}-{m:02d}"


def parse_amount(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    s = str(v).replace(',', '')
    mm = re.search(r'\d+(?:\.\d+)?', s)
    return round(float(mm.group()), 2) if mm else None


def parse_date(v, mkey):
    """v may be a datetime or text like '15 May' / '23 Dec'; mkey = 'YYYY-MM' context."""
    if isinstance(v, datetime.datetime):
        return v.strftime('%Y-%m-%d')
    if v is None:
        return None
    s = str(v).strip()
    y = int(mkey[:4])
    mm = re.match(r'^(\d{1,2})\s+([A-Za-z]+)\.?$', s)
    if mm:
        mon = MONTH_WORDS.get(mm.group(2).lower()[:3]) or MONTH_WORDS.get(mm.group(2).lower())
        if mon:
            yy = y + 1 if (mon == 1 and mkey[5:] == '12') else y
            try:
                return f"{yy}-{mon:02d}-{int(mm.group(1)):02d}"
            except ValueError:
                return None
    return None


def norm_key(name) -> str:
    s = re.sub(r'\s+', ' ', str(name).strip())
    s = re.sub(r'\s*\(.*?\)\s*$', '', s)          # strip a trailing parenthetical
    return s.lower().strip(' .,-')


def is_yes(v) -> bool:
    return str(v or '').strip().lower().startswith(('yes', 'true')) or v is True


def parse_sheet(path: str):
    wb = openpyxl.load_workbook(path, data_only=True)
    tabs = []
    for name in wb.sheetnames:
        if name in SKIP_SHEETS:
            continue
        mk = tab_month_key(name)
        if mk:
            tabs.append((mk, name))
    tabs.sort()

    raw: list[dict] = []
    warnings: list[str] = []
    for mkey, name in tabs:
        ws = wb[name]
        hdr_row, hdrs = None, []
        for r in range(1, 9):
            vals = [str(ws.cell(r, c).value or '').strip().lower() for c in range(1, ws.max_column + 1)]
            if any(v in ('client', 'client company name', 'company name') for v in vals):
                hdr_row, hdrs = r, vals
                break
        if not hdr_row:
            warnings.append(f"{name}: no header row found — skipped")
            continue

        def find(pred, last=False):
            hits = [i + 1 for i, h in enumerate(hdrs) if h and pred(h)]
            return (hits[-1] if last else hits[0]) if hits else None

        c_company = find(lambda h: h in ('client', 'client company name', 'company name'), last=True)
        c_qb = find(lambda h: h.startswith('name to search') or h == 'qb name')
        c_email = find(lambda h: 'email' in h)
        c_issue_d = find(lambda h: h in ('issue date', 'issue date '))
        c_due_d = find(lambda h: h in ('payment due date', 'due date'))
        c_amount = find(lambda h: h in ('due', '£') or (len(h) == 1 and ord(h[0]) > 127 and h not in ('$', '€')))
        c_usd = find(lambda h: h in ('$', 'usd currency?', 'usd or euro?'))
        c_eur = find(lambda h: h == '€')
        c_issued = find(lambda h: h in ('issued', 'issued?', 'invoice/ direct debit issued?'))
        c_status = find(lambda h: h == 'issue')                       # 2022-era status column
        c_am = find(lambda h: 'account manager' in h or h == 'am allocation')
        c_agreed = find(lambda h: h == 'agreed trial months')
        c_terms = find(lambda h: h == 'contract terms')
        c_mm = find(lambda h: h == 'mail merge?')
        c_pkg = next((i + 1 for i, h in enumerate(hdrs) if h == 'package' and i + 1 > (c_company or 0)), None)
        c_notes = next((i + 1 for i, h in enumerate(hdrs) if h == 'notes' and i + 1 > (c_company or 0)), None)
        c_trial = find(lambda h: h.startswith('trial month'))

        c_paid = None                                                # a cell reading 'Paid' above the header row
        for r in range(1, hdr_row):
            for c in range(1, ws.max_column + 1):
                if str(ws.cell(r, c).value or '').strip().lower() == 'paid':
                    c_paid = c
        if not c_company or not (c_amount or c_due_d or c_issue_d):
            warnings.append(f"{name}: missing key columns — skipped")
            continue

        auto = False
        for r in range(hdr_row + 1, ws.max_row + 1):
            rowvals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            joined = ' '.join(str(v) for v in rowvals if v is not None).lower()
            if 'issue automatic' in joined and not any(isinstance(v, (int, float)) for v in rowvals):
                auto = True
                continue
            comp = ws.cell(r, c_company).value
            if comp is None or str(comp).strip() == '':
                continue
            comp = re.sub(r'\s+', ' ', str(comp).strip())
            if looks_junk(comp):
                continue

            amt_raw = ws.cell(r, c_amount).value if c_amount else None
            amount = parse_amount(amt_raw) if c_amount else None
            issue_dv = ws.cell(r, c_issue_d).value if c_issue_d else None
            due_dv = ws.cell(r, c_due_d).value if c_due_d else None
            if issue_dv is None and c_status and isinstance(ws.cell(r, c_status).value, datetime.datetime):
                issue_dv = ws.cell(r, c_status).value                # 2022 era: 'Issue' sometimes holds the issue date
            issue_date = parse_date(issue_dv, mkey)
            due_date = parse_date(due_dv, mkey)
            if amount is None and not issue_date and not due_date:
                continue                                              # junk row

            issued_v = ws.cell(r, c_issued).value if c_issued else None
            if issued_v is None and c_status:
                issued_v = str(ws.cell(r, c_status).value or '').strip().lower() == 'issued'
            paid_v = ws.cell(r, c_paid).value if c_paid else None

            usd_v = ws.cell(r, c_usd).value if c_usd else None
            eur_v = ws.cell(r, c_eur).value if c_eur else None
            usd = eur = None
            if usd_v is not None:
                s = str(usd_v)
                if '€' in s:
                    eur = parse_amount(s)
                else:
                    usd = parse_amount(s)
            if eur_v is not None:
                eur = parse_amount(eur_v) or eur

            agreed = ws.cell(r, c_agreed).value if c_agreed else None
            raw.append({
                'month': mkey, 'sheet': name, 'row': r, 'company': comp, 'key': norm_key(comp),
                'qb': str(ws.cell(r, c_qb).value or '').strip() if c_qb else '',
                'email': str(ws.cell(r, c_email).value or '').strip() if c_email else '',
                'am': str(ws.cell(r, c_am).value or '').strip() if c_am else '',
                'pkg': str(ws.cell(r, c_pkg).value or '').strip() if c_pkg else '',
                'terms': str(ws.cell(r, c_terms).value or '').strip() if c_terms else '',
                'agreed': int(agreed) if isinstance(agreed, (int, float)) else None,
                'trial': ws.cell(r, c_trial).value if c_trial else None,
                'notes': str(ws.cell(r, c_notes).value or '').strip() if c_notes else '',
                'mm': is_yes(ws.cell(r, c_mm).value) if c_mm else False,
                'amount': amount if amount is not None else 0, 'amt_raw': amt_raw,
                'usd': usd, 'eur': eur,
                'issue': issue_date, 'due': due_date,
                'issued': bool(issued_v) if issued_v is not None else True,
                'paid': (bool(paid_v) if c_paid else None),           # None = this tab has no Paid column
                'auto': auto,
            })
    return raw, warnings, dict(tabs)


def sheet_totals(path: str) -> dict:
    """The sheet's own 'Total' figure per month tab, where it has one (Dec 2025 onwards) — used to check the parse."""
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for name in wb.sheetnames:
        if name in SKIP_SHEETS:
            continue
        mk = tab_month_key(name)
        if not mk:
            continue
        for row in wb[name].iter_rows(values_only=True):
            for j, c in enumerate(row):
                if isinstance(c, str) and c.strip().lower() == 'total':
                    for v in row[j + 1:j + 3]:
                        if isinstance(v, (int, float)):
                            out[mk] = round(float(v), 2)
                            break
    return out
