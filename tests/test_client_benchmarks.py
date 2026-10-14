from datetime import date

import pytest

from app import auth
from app.utils import client_benchmarks as b


def sheet(*rows):
    return [list(b.HEADERS), *rows]


def row(client='Alpha', day='01/07/2026', lead='Lead A', size='11-50'):
    return [client, day, lead, 'CEO', size, 'Software', 'Email', 'UK']


def source(*rows):
    records, quality = b.parse_rows(sheet(*rows))
    return {'records': records, 'quality': quality, 'fetched_at': '2026-10-13T12:00:00+00:00'}


def test_zeros_and_group_client_deduplication():
    data = source(row(), row('Alpha', '02/07/2026', 'Lead B', '51-200'), row('Beta', '02/09/2026', 'Lead C'))
    result = b.summarize(data, ['2026-07', '2026-08', '2026-09'])
    sme = result['groups'][1]
    assert sme['clients'] == 2
    assert sme['client_months'] == 6
    assert sme['average'] == .5  # not (2/2 + 1/1)/2 or 3 productive client-months
    assert sme['median'] == .5
    assert sme['best_month'] == 1
    assert [m['average'] for m in sme['monthly']] == [1, 0, .5]
    assert result['bands'][0]['average'] is None


def test_selected_period_defines_cohort_without_exposing_names():
    data = source(row('Old client', '01/01/2025'), row('Recent client', '02/09/2026', 'B'))
    result = b.summarize(data, ['2026-07', '2026-08', '2026-09'])
    assert result['overall']['clients'] == 1
    assert result['overall']['average'] == pytest.approx(1/3)
    assert 'Recent client' not in str(result)
    assert 'Old client' not in str(result)


def test_quality_duplicate_invalid_and_unknown_size():
    data = source(row(), row(), row(day='bad'), row(client=''), row(lead=''), row(lead='B', size='50-200'))
    assert data['quality'] == {'duplicates': 1, 'invalid_rows': 3, 'unknown_size': 1}
    result = b.summarize(data, ['2026-07'])
    assert result['overall']['leads'] == 1
    assert result['unmapped_in_period'] == 1


@pytest.mark.parametrize('raw,index', [('1-10', 0), ('11–50', 1), (' 1,001 - 5,000 ', 5), ('10001+', 7), ('10', 0), ('11', 1), ('10000', 6), ('10001', 7), ('0', None), ('50–200', None), ('unknown', None)])
def test_size_bands(raw, index):
    assert b.band_index(raw) == index


def test_date_parsing_uk_and_iso():
    assert b.parse_date('03/07/2026') == date(2026, 7, 3)
    assert b.parse_date('2026-07-03') == date(2026, 7, 3)
    assert b.parse_date('2026-06-30T23:30:00Z') == date(2026, 7, 1)
    assert b.parse_date('not a date') is None


def test_period_validation_and_year_boundary():
    assert b.default_period(date(2026, 1, 3)) == ('2025-10', '2025-12')
    assert b.period_months('2025-12', '2026-01', date(2026, 2, 1)) == ['2025-12', '2026-01']
    for start, end in [('2026-09', '2026-07'), ('x', '2026-07'), ('2026-00', '2026-07'), ('2020-01', '2026-07'), ('2026-07', '2026-10')]:
        with pytest.raises(ValueError):
            b.period_months(start, end, date(2026, 10, 13))


def test_header_banner_and_schema_failure():
    records, _ = b.parse_rows([['Results'], *sheet(row())])
    assert len(records) == 1
    with pytest.raises(b.SourceError):
        b.parse_rows([['Some other tab']])
    with pytest.raises(b.SourceError):
        b.parse_rows([list(b.HEADERS) + ['Client']])


def test_page_api_and_reporting_navigation(client, monkeypatch):
    monkeypatch.setattr(b, 'load_source', lambda: source(row()))
    assert client.get('/client-benchmarks').status_code == 200
    assert 'href="/client-benchmarks"' in client.get('/').text
    response = client.get('/api/client-benchmarks?start=2026-07&end=2026-09')
    assert response.status_code == 200
    assert response.json()['overall']['average'] == pytest.approx(1/3)
    assert response.headers['cache-control'] == 'no-store'
    assert client.get('/api/client-benchmarks?start=invalid').status_code == 400


def test_access_control_covers_api_and_page(client, monkeypatch):
    monkeypatch.delenv('AUTH_DISABLED', raising=False)
    monkeypatch.setattr(auth, 'get_session_user', lambda request: None)
    assert client.get('/client-benchmarks', follow_redirects=False).status_code == 303
    assert client.get('/api/client-benchmarks').status_code == 401
    monkeypatch.setattr(auth, 'get_session_user', lambda request: {'role': 'viewer', 'tools': []})
    assert client.get('/client-benchmarks').status_code == 403
    assert client.get('/api/client-benchmarks').status_code == 403
    assert 'href="/client-benchmarks"' not in client.get('/').text
    monkeypatch.setattr(auth, 'get_session_user', lambda request: {'role': 'sdr', 'tools': ['client_benchmarks']})
    assert 'href="/client-benchmarks"' in client.get('/').text


def test_source_failure_does_not_substitute_sample_data(client, monkeypatch):
    def fail():
        raise b.SourceError('Sheet unavailable')
    monkeypatch.setattr(b, 'load_source', fail)
    response = client.get('/api/client-benchmarks')
    assert response.status_code == 503
    assert response.json() == {'error': 'Sheet unavailable'}


def test_public_source_cache_and_html_login_rejection(monkeypatch):
    monkeypatch.setattr(b, '_cache', None)
    monkeypatch.setattr(b, '_is_configured', lambda: False)
    calls = []
    class Response:
        headers = {'Content-Type': 'text/csv'}
        content = (','.join(b.HEADERS) + '\n' + ','.join(row())).encode()
        def raise_for_status(self):
            pass
    def get(url, **kwargs):
        calls.append(url)
        assert kwargs['params']['gid'] == b.GID
        assert kwargs['timeout'] == 20
        return Response()
    monkeypatch.setattr(b.requests, 'get', get)
    assert len(b.load_source()['records']) == 1
    assert b.load_source()['records'] == b.load_source()['records']
    assert len(calls) == 1
    monkeypatch.setattr(b, '_cache', None)
    Response.headers = {'Content-Type': 'text/html'}
    with pytest.raises(b.SourceError, match='requires access'):
        b.load_source()


def test_authenticated_source_selects_gid_and_only_reads(monkeypatch):
    monkeypatch.setattr(b, '_cache', None)
    monkeypatch.setattr(b, '_is_configured', lambda: True)
    class Response:
        def __init__(self, payload): self.payload = payload
        def raise_for_status(self): pass
        def json(self): return self.payload
    class Session:
        def get(self, url, **kwargs):
            if '/values/' in url:
                assert '%27Client%20leads%27%21A%3AZ' in url
                return Response({'values': sheet(row())})
            return Response({'sheets': [{'properties': {'sheetId': 1, 'title': 'Wrong'}}, {'properties': {'sheetId': b.GID, 'title': 'Client leads'}}]})
    monkeypatch.setattr(b, '_read_session', lambda: Session())
    assert len(b.load_source()['records']) == 1


def test_auth_failure_is_sanitized(monkeypatch):
    monkeypatch.setattr(b, '_cache', None)
    monkeypatch.setattr(b, '_is_configured', lambda: True)
    def fail(): raise RuntimeError('sensitive credential content')
    monkeypatch.setattr(b, '_read_session', fail)
    with pytest.raises(b.SourceError) as error:
        b.load_source()
    assert 'sensitive' not in str(error.value)


def excel_bytes(tabs):
    from io import BytesIO
    import openpyxl
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in tabs.items():
        sheet = workbook.create_sheet(name)
        for record in rows:
            sheet.append(record)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_excel_prefers_current_raw_tab_and_counts_unnamed_leads():
    from app.utils.benchmark_upload import read_excel
    from datetime import datetime
    headers = ['Timestamp', 'Client', 'Lead company size', 'Lead category']
    raw = [headers,
           [datetime(2026, 7, 2, 10, 30), 'Alpha', '11-50', 'Lead'],
           [datetime(2026, 7, 2, 10, 30), 'Alpha', '11-50', 'Lead'],
           [datetime(2026, 7, 3), 'Alpha', '11-50', 'Interested'],
           [datetime(2026, 7, 3), 'Unstuck', '11-50', 'Lead'],
           [datetime(2026, 7, 4), 'Beta', 'unknown', 'Lead']]
    s = read_excel(excel_bytes({'Responses (clean)': sheet(row()), 'Form Responses': raw}))
    assert s['source_tab'] == 'Form Responses'
    assert len(s['records']) == 3
    assert s['quality']['duplicates'] == 1
    assert s['quality']['excluded_category'] == 1
    assert any(r['client'] == 'unstuck' for r in s['records'])
    assert s['quality']['unknown_size'] == 1
    assert s['date_max'] == '2026-07-04'


def test_excel_standard_headers_and_empty_workbook():
    from app.utils.benchmark_upload import read_excel
    assert len(read_excel(excel_bytes({'Leads': sheet(row())}))['records']) == 1
    with pytest.raises(b.SourceError, match='No lead data tab'):
        read_excel(excel_bytes({'Other': [['Unknown']]}))
    with pytest.raises(b.SourceError, match='Multiple lead data tabs'):
        read_excel(excel_bytes({'A': sheet(row()), 'B': sheet(row())}))
    with pytest.raises(b.SourceError, match='valid, unencrypted'):
        read_excel(b'not an xlsx')


def test_excel_upload_api_and_permission(client, monkeypatch):
    content = excel_bytes({'Leads': sheet(row())})
    result = client.post('/api/client-benchmarks/upload', files={'file': ('leads.xlsx', content)}, data={'start': '2026-07', 'end': '2026-09'})
    assert result.status_code == 200
    assert result.json()['source_kind'] == 'excel'
    assert result.json()['overall']['average'] == pytest.approx(1/3)
    assert result.headers['cache-control'] == 'no-store'
    assert client.post('/api/client-benchmarks/upload', files={'file': ('leads.csv', content)}).status_code == 400
    assert client.post('/api/client-benchmarks/upload', files={'file': ('leads.xlsx', content)}, data={'start': 'bad'}).status_code == 400
    monkeypatch.delenv('AUTH_DISABLED', raising=False)
    monkeypatch.setattr(auth, 'get_session_user', lambda request: {'role': 'viewer', 'tools': []})
    assert client.post('/api/client-benchmarks/upload', files={'file': ('leads.xlsx', content)}).status_code == 403


def test_excel_limits_and_internal_client(monkeypatch):
    from app.utils import benchmark_upload as u
    monkeypatch.setattr(u, 'MAX_UPLOAD_BYTES', 5)
    with pytest.raises(b.SourceError, match='20 MB'):
        u.read_excel(b'123456')
    assert source(row(client=' Unstuck '))['records'][0]['client'] == 'unstuck'


ACTIVITY = ['Campaign ID', 'Client', 'Month', 'Company size', 'Active', 'Leads']


def test_active_campaign_months_include_zero_leads_and_dedup_groups():
    from app.utils.benchmark_upload import read_excel
    rows = [ACTIVITY,
            ['c1', 'Alpha', '2026-07', '11-50', 'Yes', 4],
            ['c1', 'Alpha', '2026-07', '51-200', 'Yes', 2],
            ['c1', 'Alpha', '2026-08', '11-50', 'Yes', 0],
            ['c2', 'Unstuck', '2026-07', '11-50', True, 3],
            ['c2', 'Unstuck', '2026-08', '11-50', False, 0],
            ['c3', 'Beta', '2026-08', '51-200', 1, 0]]
    s = read_excel(excel_bytes({'Activity': rows, 'Old raw records': sheet(row())}))
    r = b.summarize(s, ['2026-07', '2026-08', '2026-09'])
    assert r['basis'] == 'active_campaign_month'
    sme = r['groups'][1]
    assert sme['leads'] == 9
    assert sme['exposure_months'] == 4  # overlapping c1/July counts once
    assert sme['average'] == 2.25  # zero-lead c1/August and c3/August count
    assert sme['campaigns'] == 3
    assert sme['clients'] == 3
    assert sme['unstuck_leads'] == 3
    assert sme['median'] == 3
    assert [m['average'] for m in sme['monthly']] == [4.5, 0, None]
    assert s['source_tab'] == 'Activity'


@pytest.mark.parametrize('bad', [
    ['c1','Alpha','2026-07','11-50','Yes',None],
    ['c1','Alpha','2026-07','11-50','Yes',-1],
    ['c1','Alpha','2026-07','11-50','Yes',1.5],
    ['c1','Alpha','2026-07','11-50','No',1],
    ['c1','Alpha','2026-07','11-50',None,0],
    ['c1','Alpha','2026-13','11-50','Yes',0],
    ['c1','Alpha','2026-07','50-200','Yes',0],
])
def test_activity_rejects_ambiguous_inputs(bad):
    from app.utils.benchmark_upload import read_excel
    with pytest.raises(b.SourceError):
        read_excel(excel_bytes({'Activity': [ACTIVITY, bad]}))


def test_activity_duplicate_and_conflicting_owners_rejected():
    from app.utils.benchmark_upload import read_excel
    base = ['c1','Alpha','2026-07','11-50','Yes',1]
    for other in [base, ['c1','Beta','2026-08','11-50','Yes',1]]:
        with pytest.raises(b.SourceError):
            read_excel(excel_bytes({'Activity': [ACTIVITY, base, other]}))


def test_activity_zero_only_campaign_and_template_endpoint(client):
    from app.utils.benchmark_upload import read_excel
    s = read_excel(excel_bytes({'Activity': [ACTIVITY, ['c1','Alpha','2026-07','11-50','Yes',0]]}))
    r = b.summarize(s, ['2026-07'])
    assert r['overall']['average'] == 0
    assert r['overall']['exposure_months'] == 1
    result = client.get('/api/client-benchmarks/activity-template')
    assert result.status_code == 200
    assert 'spreadsheetml' in result.headers['content-type']
    from io import BytesIO
    import openpyxl
    w = openpyxl.load_workbook(BytesIO(result.content))
    assert list(next(w['Campaign activity'].values)) == ACTIVITY
    assert w['Campaign activity'].max_row == 1


def test_calendar_source_never_claims_campaign_activity():
    r = b.summarize(source(row(client='Unstuck')), ['2026-07','2026-08'])
    assert r['basis'] == 'calendar_account_month'
    assert r['overall']['unstuck_leads'] == 1
    assert r['overall']['average'] == .5


def test_delivery_month_proxy_excludes_untracked_clients_and_uses_full_history():
    from app.utils.benchmark_upload import read_excel
    book = excel_bytes({'Leads': sheet(row('Alpha','01/07/2026','A'), row('Alpha','01/09/2026','B'), row('Unstuck','01/09/2026','C')),
                       'Delivery months': [['Client','Lead sheet client','Delivery months'], ['Alpha Billing','Alpha',2],['Untracked','Untracked',20]]})
    s = read_excel(book)
    r = b.summarize(s, ['2026-09'])
    assert r['basis'] == 'delivery_month'
    assert r['overall']['leads'] == 2  # full history, not date-filtered
    assert r['overall']['exposure_months'] == 2  # not 22
    assert r['overall']['average'] == 1
    assert r['matched_accounts'] == 1
    assert r['unmatched_accounts'] == 1
    assert r['unmatched_delivery_months'] == 20
    assert not r['unstuck_in_roster']
    assert r['overall']['monthly'] == []


def test_delivery_month_proxy_deduplicates_client_in_group():
    from app.utils.benchmark_upload import read_excel
    s = read_excel(excel_bytes({'Leads': sheet(row(),row(lead='B',size='51-200')),
        'Delivery months': [['Client','Lead sheet client','Delivery months'],['Alpha','Alpha',4]]}))
    r = b.summarize(s, ['2026-07'])
    assert r['groups'][1]['exposure_months'] == 4
    assert r['groups'][1]['average'] == .5


@pytest.mark.parametrize('count', [0,-1,1.5,'nan','bad'])
def test_invalid_delivery_counts_rejected(count):
    from app.utils.benchmark_upload import read_excel
    with pytest.raises(b.SourceError):
        read_excel(excel_bytes({'Leads':sheet(row()),'Delivery months':[['Client','Lead sheet client','Delivery months'],['Alpha','Alpha',count]]}))


def test_selected_examples_combine_bands_keep_months_separate_and_hide_names():
    data = source(row('Confidential', '01/07/2026', 'A'),
                  row('Confidential', '02/07/2026', 'B', '51-200'),
                  row('Confidential', '01/08/2026', 'C'),
                  row('Confidential', '02/08/2026', 'D'),
                  row('Other', '01/09/2026', 'E'))
    result = b.summarize(data, ['2026-07', '2026-08'])
    example = result['examples'][1]
    assert example['leads'] == 2
    assert example['month'] == '2026-08'
    assert example['candidate_accounts'] == 1
    assert example['observed_account_months'] == 2
    assert 'Confidential' not in str(result['examples'])
    assert result['examples'][0]['leads'] is None
    assert result['overall']['average'] == 2


def test_billing_examples_have_separate_window_and_unstuck_disclosure(monkeypatch):
    from app.utils.benchmark_upload import read_excel
    monkeypatch.setattr(b, 'london_today', lambda: date(2026, 10, 13))
    data = read_excel(excel_bytes({'Leads': sheet(
        row('Alpha', '01/01/2024', 'Old'), row('Alpha', '01/09/2026', 'A'),
        row('Unstuck', '01/09/2026', 'B'), row('Unstuck', '02/09/2026', 'C'),
        row('Unstuck', '01/10/2026', 'Current'), row('Outside roster', '01/09/2026', 'D')),
        'Delivery months': [['Client','Lead sheet client','Delivery months'], ['Alpha','Alpha',2]]}))
    result = b.summarize(data, ['2026-09'])
    assert result['overall']['average'] == 1
    assert result['overall']['leads'] == 2
    assert result['example_window'][0] == '2025-10'
    assert result['example_window'][-1] == '2026-09'
    example = result['examples'][1]
    assert example['leads'] == 2
    assert example['is_unstuck']
    assert example['account_label'] == 'Unstuck — own prospecting'
    assert example['candidate_accounts'] == 2


def test_zero_activity_does_not_fabricate_strong_example():
    from app.utils.benchmark_upload import read_excel
    data = read_excel(excel_bytes({'Activity': [ACTIVITY, ['c1','Alpha','2026-07','11-50','Yes',0]]}))
    assert b.summarize(data, ['2026-07'])['examples'][1]['leads'] is None


def test_benchmark_google_session_is_separate_and_read_only(monkeypatch):
    import base64
    import json
    from app.utils import google_sheets
    original_session = google_sheets._session_cache
    original_client = google_sheets._gc_cache
    monkeypatch.setenv('GOOGLE_SHEETS_SA_JSON', base64.b64encode(json.dumps({'client_email': 'synthetic@example.invalid'}).encode()).decode())
    sentinel = object()
    def credentials(info, scopes):
        assert info == {'client_email': 'synthetic@example.invalid'}
        assert scopes == ['https://www.googleapis.com/auth/spreadsheets.readonly']
        return sentinel
    monkeypatch.setattr(b.Credentials, 'from_service_account_info', credentials)
    monkeypatch.setattr(b, 'AuthorizedSession', lambda credentials: ('independent', credentials))
    assert b._read_session() == ('independent', sentinel)
    assert google_sheets._session_cache is original_session
    assert google_sheets._gc_cache is original_client


def test_benchmark_failure_leaves_roi_and_home_available(client, monkeypatch):
    def fail():
        raise b.SourceError('Unavailable')
    monkeypatch.setattr(b, 'load_source', fail)
    assert client.get('/api/client-benchmarks').status_code == 503
    for path in ['/', '/roi-calculator']:
        response = client.get(path)
        assert response.status_code == 200
        assert '/static/client-benchmarks.css' not in response.text
        assert '/static/client-benchmarks.js' not in response.text
